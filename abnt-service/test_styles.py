import io
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from models import DocumentPayload, CoverData, SectionData, ReferenceData
from styles import build_styled_document, ref_vancouver, ref_apa, autor_vancouver, autor_apa, partes_do_nome
from main import app
from fastapi.testclient import TestClient

REFS = [
    ReferenceData(authors=["SOUZA, Carla", "Lima, Pedro Henrique"], title="Água e plantas C4", year=2019, venue="Revista Brasileira de Botânica", doi="10.1000/rbb.2019.1"),
    ReferenceData(authors=["Ana Beatriz Silva"], title="Fisiologia vegetal", year=2021),
]


def _build(style, **kw):
    p = DocumentPayload(
        title="Fotossíntese C4", work_type="Artigo", style=style,
        cover=CoverData(author="Fulano", institution="UFC", course="Biologia", city="Fortaleza", year="2026"),
        resumo="Resumo do trabalho.", palavras_chave=["fotossíntese", "água"],
        sections=[
            SectionData(type="introducao", title="1 INTRODUÇÃO", content="Texto com **destaque** e citação [1]."),
            SectionData(type="referencias", title="REFERÊNCIAS", content="SOUZA, C. Texto ABNT antigo."),
        ],
        references=kw.pop("references", REFS), format="docx", **kw,
    )
    return Document(io.BytesIO(build_styled_document(p)))


def _texto(doc):
    return [p.text for p in doc.paragraphs]


def test_nomes():
    assert partes_do_nome("SILVA, Ana Beatriz") == ("Silva", ["Ana", "Beatriz"])
    assert autor_vancouver("Ana Beatriz da Silva") == "Silva AB"
    assert autor_apa("SOUZA, Carla") == "Souza, C."


def test_vancouver_numerada_na_ordem_e_formato_icmje():
    doc = _build("vancouver")
    t = _texto(doc)
    i = t.index("Referências")
    assert t[i + 1] == "1. Souza C, Lima PH. Água e plantas C4. Revista Brasileira de Botânica. 2019. doi:10.1000/rbb.2019.1"
    assert t[i + 2] == "2. Silva AB. Fisiologia vegetal. 2021."
    # a seção REFERÊNCIAS do texto (ABNT) não vira corpo; título sem número
    assert "SOUZA, C. Texto ABNT antigo." not in t
    assert "Introdução" in t and "1 INTRODUÇÃO" not in t
    assert round(doc.sections[0].left_margin.cm, 1) == 2.5
    assert "Descritores: fotossíntese; água." in t


def test_apa_alfabetica_recuo_deslocado_italico_e_espaco_duplo():
    doc = _build("apa")
    ps = doc.paragraphs
    t = _texto(doc)
    i = t.index("Referências")
    assert t[i + 1] == "Silva, A. B. (2021). Fisiologia vegetal."
    assert t[i + 2] == "Souza, C., & Lima, P. H. (2019). Água e plantas C4. Revista Brasileira de Botânica. https://doi.org/10.1000/rbb.2019.1"
    ref = ps[i + 2]
    assert round(ref.paragraph_format.first_line_indent.cm, 2) == -1.27
    assert [r.text for r in ref.runs if r.italic] == ["Revista Brasileira de Botânica"]
    assert [r.text for r in ps[i + 1].runs if r.italic] == ["Fisiologia vegetal"]
    corpo = next(p for p in ps if "destaque" in p.text)
    assert corpo.paragraph_format.line_spacing == 2.0
    assert corpo.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert [r.text for r in corpo.runs if r.bold] == ["destaque"]
    titulo = next(p for p in ps if p.text == "Introdução")
    assert titulo.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_referencia_so_com_texto_sai_como_veio():
    r = ReferenceData(text="Documento sem campos estruturados.")
    assert ref_vancouver(r) == "Documento sem campos estruturados."
    assert ref_apa(r) == [("Documento sem campos estruturados.", False)]


def test_rota_format_despacha_pela_norma_e_nomeia_o_arquivo():
    c = TestClient(app)
    body = {
        "title": "Teste", "work_type": "Artigo", "style": "apa", "format": "docx",
        "cover": {"author": "A", "institution": "", "course": "", "city": "", "year": "2026"},
        "resumo": "r", "palavras_chave": ["x"], "sections": [{"type": "introducao", "title": "Intro", "content": "c", "subsections": []}],
    }
    r = c.post("/format", json=body)
    assert r.status_code == 200 and "_APA.docx" in r.headers["content-disposition"]
    body["style"] = "abnt"; body["cover"]["institution"] = "UFC"; body["cover"]["city"] = "Fortaleza"
    r = c.post("/format", json=body)
    assert r.status_code == 200 and "_ABNT.docx" in r.headers["content-disposition"]
